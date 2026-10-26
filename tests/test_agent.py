import asyncio
import logging
from io import BytesIO

from docx import Document
from fakes import APPROVED, PLAN, SECTION, TABLE, happy_path

import llm
import main

REQUEST = {"request": "Write minutes for our sprint planning meeting today."}


def headings(docx_bytes: bytes) -> list[str]:
    doc = Document(BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]


def test_agent_plans_writes_and_returns_a_document(client, fake_llm):
    response = client.post("/agent", json=REQUEST)

    assert response.status_code == 200
    body = response.json()
    assert body["summary"] == (
        "Created 'Sprint Planning Minutes' with 2 sections from 3 tasks. "
        "Flagged 1 assumption. Review approved after 0 revisions."
    )
    assert [step for step, _ in fake_llm.calls] == [
        "plan",
        "section",
        "table",
        "review",
    ]

    download = client.get(body["download_url"])
    assert download.status_code == 200
    doc = Document(BytesIO(download.content))
    # The numeric cell from the model was coerced to text.
    assert doc.tables[0].rows[1].cells[2].text == "2"
    assert headings(download.content) == [
        "Assumptions & Risks",
        "Decisions",
        "Action Items",
    ]


def test_review_rewrites_in_place_and_stops_after_one_revision(client, fake_llm):
    async def never_satisfied(step, model, user):
        if step == "review":
            return {
                "approved": False,
                "gaps": ["Owner differs between sections"],
                # write_section on purpose: the planned tool (make_table) must win.
                "fix_tasks": [
                    {
                        "title": "Action Items",
                        "tool": "write_section",
                        "details": "Fix owner",
                    }
                ],
            }
        return await happy_path(step, model, user)

    fake_llm.responder = never_satisfied
    response = client.post("/agent", json=REQUEST)

    body = response.json()
    assert body["revision_count"] == 1
    assert [step for step, _ in fake_llm.calls].count("review") == 2
    assert body["tasks"][-1] == {
        "title": "Action Items",
        "tool": "make_table",
        "details": "Fix owner",
        "source": "review",
    }
    content = client.get(body["download_url"]).content
    assert headings(content) == ["Assumptions & Risks", "Decisions", "Action Items"]
    assert len(Document(BytesIO(content)).tables) == 1


def test_falls_back_when_primary_model_returns_invalid_json(client, fake_llm):
    async def broken_primary(step, model, user):
        if model == llm.PRIMARY_MODEL:
            return '{"title": "missing everything else"}'
        return await happy_path(step, model, user)

    fake_llm.responder = broken_primary
    response = client.post("/agent", json=REQUEST)

    assert response.status_code == 200
    assert fake_llm.calls[:2] == [
        ("plan", llm.PRIMARY_MODEL),
        ("plan", llm.FALLBACK_MODEL),
    ]


def test_returns_503_when_every_model_fails(client, fake_llm):
    async def always_broken(step, model, user):
        return "not json"

    fake_llm.responder = always_broken
    response = client.post("/agent", json=REQUEST)

    assert response.status_code == 503


def test_returns_504_when_the_agent_runs_past_its_deadline(
    client, fake_llm, monkeypatch
):
    async def slow(step, model, user):
        await asyncio.sleep(1)
        return await happy_path(step, model, user)

    fake_llm.responder = slow
    monkeypatch.setattr(main, "AGENT_TIMEOUT_SECONDS", 0.1)
    response = client.post("/agent", json=REQUEST)

    assert response.status_code == 504


def test_unknown_tool_runs_as_write_section(client, fake_llm):
    plan = {
        **PLAN,
        "tasks": [{"title": "Market Research", "tool": "search_web", "details": "x"}],
    }

    async def unknown_tool(step, model, user):
        return plan if step == "plan" else await happy_path(step, model, user)

    fake_llm.responder = unknown_tool
    response = client.post("/agent", json=REQUEST)

    assert response.json()["tasks"][0]["tool"] == "write_section"


def test_every_log_line_of_a_request_carries_its_request_id(client):
    records: list[logging.LogRecord] = []
    handler = logging.Handler()
    handler.emit = records.append
    handler.addFilter(main.RequestIdFilter())
    logging.getLogger().addHandler(handler)
    try:
        first = client.post("/agent", json=REQUEST).headers["X-Request-ID"]
        second = client.post("/agent", json=REQUEST).headers["X-Request-ID"]
    finally:
        logging.getLogger().removeHandler(handler)

    agent_ids = {
        r.request_id for r in records if r.name in ("planner", "executor", "reviewer")
    }
    assert first != second
    assert agent_ids == {first, second}


def test_rejects_requests_that_are_too_short(client, fake_llm):
    response = client.post("/agent", json={"request": "hi"})

    assert response.status_code == 422
    assert fake_llm.calls == []


def test_document_download_validates_the_id(client):
    assert client.get("/documents/not-a-real-id").status_code == 422
    assert client.get("/documents/" + "0" * 32).status_code == 404


def test_fixtures_match_the_models():
    # Guards the fake replies against drifting from the real schemas.
    from models import Review
    from planner import Plan
    from tools import SectionText, TableRows

    Plan.model_validate(PLAN)
    SectionText.model_validate(SECTION)
    TableRows.model_validate(TABLE)
    Review.model_validate(APPROVED)
