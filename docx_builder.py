from datetime import date
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.document import Document as DocxDocument

from models import AgentState

OUTPUT_DIR = Path(__file__).parent / "output"


def add_table(doc: DocxDocument, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def add_text(doc: DocxDocument, text: str) -> None:
    for line in text.replace("**", "").splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith(("- ", "• ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)


def build_docx(state: AgentState) -> AgentState:
    OUTPUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    doc.add_heading(state.title or "Document", level=0)
    doc.add_paragraph(f"{state.doc_type} · Generated {date.today():%d %B %Y}")

    if state.assumptions:
        doc.add_heading("Assumptions & Risks", level=1)
        for assumption in state.assumptions:
            doc.add_paragraph(assumption, style="List Bullet")

    for section in state.sections:
        doc.add_heading(section.heading, level=1)
        if section.text:
            add_text(doc, section.text)
        if section.table:
            add_table(doc, section.table)
            
    state.document_id = uuid4().hex
    doc.save(OUTPUT_DIR / f"{state.document_id}.docx")
    return state
